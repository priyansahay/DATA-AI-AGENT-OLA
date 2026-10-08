from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "UBER_AI_AGENT_CODE_ANALYSIS.docx"
INK = RGBColor(36, 48, 58)
TEAL = RGBColor(0, 112, 116)
MUTED = RGBColor(93, 107, 117)


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        from docx.oxml import OxmlElement

        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=RGBColor(255, 255, 255))
        set_cell_shading(table.rows[0].cells[index], "006F73")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], str(value))
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F0F5F5")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_numbered(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_callout(document, label, text, fill="E8F2F1"):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title = paragraph.add_run(f"{label}  ")
    title.bold = True
    title.font.color.rgb = TEAL
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color in (("Title", 30, INK), ("Heading 1", 19, INK), ("Heading 2", 13, TEAL), ("Heading 3", 10.5, INK)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("UBER AI AGENT  /  CODE REVIEW")
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Project analysis  |  08 October 2026")
    footer_run.font.name = "Aptos"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED


def build_report():
    document = Document()
    configure(document)

    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(34)
    eyebrow.paragraph_format.space_after = Pt(8)
    run = eyebrow.add_run("ENGINEERING REVIEW  /  01")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = TEAL

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(7)
    title.add_run("UBER AI AGENT")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Codebase analysis, architecture, workflows, and risk review")
    run.font.size = Pt(15)
    run.font.color.rgb = MUTED

    add_callout(
        document,
        "Scope",
        "Review of the Python source, LangGraph diagrams, PostgreSQL loading workflow, requirements, and checked-in sample data. Prepared 08 October 2026.",
    )

    add_heading(document, "Executive Summary")
    document.add_paragraph(
        "This repository is a prototype with two natural-language data agents. The ETL agent uses LangGraph and model-selected tools to extract API data or transform a local file. The SQL analyst converts a question into PostgreSQL, asks an LLM judge whether it appears read-only, executes it, and formats the result. The project also provides synthetic rideshare CSV data and a script that creates and loads a PostgreSQL schema."
    )
    document.add_paragraph(
        "The core workflows are visible and the included graph images match the current graph definitions. The main production blockers are security and lifecycle concerns: LLM-generated Python is executed with exec, SQL safety depends on an LLM judgment, database setup scripts truncate tables, and database utility code contains a hard-coded credential and performs work during import. These should be addressed before exposing the agents to untrusted users or valuable data."
    )

    add_heading(document, "Repository at a Glance")
    add_table(
        document,
        ["Area", "Files", "Responsibility"],
        [
            ("Agents", "agents/etl_analyst.py", "Tool-using ETL graph and runnable example."),
            ("Agents", "agents/sql_analyst.py", "Natural-language-to-SQL graph with a safety branch."),
            ("Models", "Models/schema.py", "Pydantic state contracts for SQL and ETL graphs."),
            ("Utilities", "utils/etl_tools.py", "API extraction, local path handling, file preview, code execution."),
            ("Utilities", "utils/database.py", "PostgreSQL metadata inspection and query execution."),
            ("Utilities", "utils/llmpick.py", "Selects configured OpenAI or Anthropic chat models."),
            ("Data", "data/*.csv", "Synthetic users, vehicles, rides, payments, and ratings."),
            ("Database", "feed_db.py", "Creates tables/indexes and reloads CSV records."),
            ("Entry point", "main.py", "Currently empty; no unified application entry point."),
        ],
        [1.0, 2.0, 4.0],
    )

    add_heading(document, "1. ETL Agent")
    document.add_paragraph(
        "The ETL graph accepts a chat message. Its LLM node sends the system instructions and conversation to a model bound to two tools. When the assistant returns tool calls, the router sends execution to the tool node; each result is added as a ToolMessage and control loops back to the LLM. A response without tool calls ends the graph. Nodes return only newly created messages, which is important because ETLAgentSchema declares an additive message reducer."
    )
    graph_path = ROOT / "etl_analyst_graph.png"
    if graph_path.exists():
        document.add_picture(str(graph_path), width=Inches(2.3))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph("Figure 1. ETL agent graph: LLM routing, tool execution loop, and completion path.")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(8)
        caption.runs[0].font.color.rgb = MUTED

    add_heading(document, "ETL tools and example", 2)
    add_bullet(document, "extract_load_tool calls an HTTP endpoint, expects JSON with a results collection, normalizes that collection, and writes CSV, JSON Lines, or Parquet under the project path.")
    add_bullet(document, "transform_load_tool resolves the input/output paths, previews up to three rows, requests Pandas transformation code from an LLM, then executes the returned code.")
    add_bullet(document, "The active example reads data/extract/extracted_data.csv and asks for a Bulbasaur-only CSV in data/transform. The checked-in result contains one row.")
    add_callout(
        document,
        "Important limitation",
        "transform_load_context currently uses the input file to produce a preview; output_folder and output_format do not themselves perform a transformation or write a file. The generated code is responsible for writing the output.",
        "FFF3DF",
    )

    add_heading(document, "2. SQL Analyst")
    document.add_paragraph(
        "The SQL graph curates the question, fetches public-schema table/column/sample details, prompts a model for SQL, and asks a structured-output judge for Yes/No. Yes routes to query execution and answer generation; No returns a cancellation response. The graph image visualizes both terminal paths."
    )
    graph_path = ROOT / "sql_analyst_graph.png"
    if graph_path.exists():
        document.add_picture(str(graph_path), width=Inches(3.2))
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption = document.add_paragraph("Figure 2. SQL analyst graph: question preparation, SQL generation, safety gate, and response paths.")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(8)
        caption.runs[0].font.color.rgb = MUTED

    add_heading(document, "SQL workflow and state", 2)
    add_numbered(document, "curate_ques rewrites the user question and adds it to the messages state.")
    add_numbered(document, "prompt_query_context obtains schema/sample data from PostgreSQL and builds the generation prompt.")
    add_numbered(document, "generate_sql produces a query; is_safe_sql requests a structured Yes/No decision and comments.")
    add_numbered(document, "The conditional edge either cancels or executes the query, then represent_final_answer summarizes returned rows.")
    add_bullet(document, "SQL_ANALYST_NODE_RESPONSE.txt captures messages, generated SQL, execution results, and the schema prompt from the runnable example.")
    add_bullet(document, "The current AgentSchema uses an additive reducer for messages, while several SQL nodes mutate and return the whole Pydantic state. The checked-in response text contains repeated copies of the curated HumanMessage, consistent with duplicate accumulation; return message deltas or use a reducer-safe update pattern.")

    add_heading(document, "3. Database and Data")
    document.add_paragraph(
        "feed_db.py defines the public PostgreSQL schema, indexes, and CSV load order. Foreign keys connect vehicles to users, rides to riders/drivers, payments to rides/users, and ratings to rides/riders/drivers. The script loads users before dependent tables and prints final record counts."
    )
    add_table(
        document,
        ["Table / file", "Rows", "Key content"],
        [
            ("users", "10,000", "Riders/drivers, contact details, province, signup and active status."),
            ("vehicles", "3,000", "Driver, make/model, year, plate, color, active status."),
            ("rides", "20,000", "Participants, times, coordinates, distance, fare, status and cancellation."),
            ("payments", "16,073", "Ride/user, amount, method/status, transaction and payment time."),
            ("ratings", "12,000", "Ride, rider/driver, 1-5 rating, comment and timestamp."),
            ("Extracted Pokémon", "20", "name and URL from the API results."),
            ("Transformed Pokémon", "1", "Bulbasaur-only output sample."),
        ],
        [1.65, 0.75, 5.6],
    )

    add_heading(document, "4. Configuration and How to Run")
    document.add_paragraph("The code loads environment values from a local .env file. Do not commit that file or include its values in reports/logs. The SQL agent expects host, port, database, user, and password. The database utility also has a separate hard-coded connection configuration in source that should be removed.")
    add_table(
        document,
        ["Purpose", "Command / configuration"],
        [
            ("Install declared dependencies", "python3 -m pip install -r requirements.txt"),
            ("Load/rebuild sample database", "python3 feed_db.py  (destructive: truncates existing tables first)"),
            ("Run ETL demonstration", "python3 agents/etl_analyst.py"),
            ("Run SQL analyst demonstration", "python3 agents/sql_analyst.py"),
            ("Database settings", ".env keys: host, port, database, user, password"),
        ],
        [2.1, 5.9],
    )
    document.add_paragraph(
        "requirements.txt lists LangChain, LangGraph, provider integrations, Pydantic, psycopg2-binary, and IPython. Source also directly imports python-dotenv, requests, and pandas; these should be explicit dependencies. Parquet output may additionally require pyarrow or another Pandas Parquet engine."
    )

    add_heading(document, "5. Findings and Recommended Priorities")
    add_table(
        document,
        ["Priority", "Finding", "Recommended action"],
        [
            ("Critical", "LLM-generated Python is passed to exec in utils/etl_tools.py. A prompt or model response can run arbitrary code with the process account's permissions.", "Remove arbitrary execution; use a constrained transform DSL/allowlisted operations. If execution is unavoidable, isolate it in a locked-down container with no secrets/network and strict resource limits."),
            ("Critical", "utils/database.py contains a hard-coded database credential and creates a connection/query at module import time.", "Remove and rotate the credential; load configuration only from environment/secret storage. Move all connection/query/file-write work into explicit functions and close resources with context managers."),
            ("High", "SQL is generated by an LLM and guarded by another LLM before direct execution. The judge is not a deterministic SQL security boundary.", "Use a read-only DB role and read-only transaction; parse SQL and allow only a single SELECT/approved statements. Apply statement timeout and row limits in the database."),
            ("High", "feed_db.py truncates all five tables with CASCADE before loading CSVs, and this runs at script top level.", "Make destructive reset opt-in, require an explicit flag/confirmation, and isolate schema setup from routine ingestion."),
            ("High", "Database configuration and side effects are spread across SQL nodes and DatabaseUtil; importing the utility performs a live DB operation and writes test_schema.txt.", "Make DatabaseUtil construction side-effect free; move diagnostics to an explicit command and centralize connection configuration."),
            ("Medium", "SQL state messages use an additive reducer while nodes mutate/return the full state, producing duplicated messages in the saved response artifact.", "Return only changed message deltas or switch to a consistent reducer-managed state-update style; add a graph-cycle regression test."),
            ("Medium", "ETL API extraction has no request timeout and assumes every response has a results key; error handling mainly covers HTTP transport failures.", "Add connect/read timeouts, validate response shape, handle JSON/schema failures, and report actionable errors."),
            ("Medium", "requirements.txt omits directly imported runtime packages, and no automated tests are present in the project inventory.", "Declare direct dependencies and add unit tests for path resolution, tool-call message ordering, ETL formats, SQL routing, and destructive-load safeguards."),
        ],
        [0.8, 3.25, 4.0],
    )

    add_heading(document, "6. Overall Assessment")
    document.add_paragraph(
        "The repository is a useful learning prototype: graph definitions are small, the two workflows are understandable, and the included sample data and graph images make behavior inspectable. It is not yet suitable for production use because execution and SQL controls are prompt-based, destructive database actions are implicit, and importing a utility can access the database. Address the Critical findings first, then stabilize state updates and add automated tests before expanding the agents."
    )
    add_callout(
        document,
        "Evidence note",
        "This report is based on the workspace source and artifacts available on 08 October 2026. It does not read or reproduce .env secrets, contact external services, or validate the live PostgreSQL instance.",
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()