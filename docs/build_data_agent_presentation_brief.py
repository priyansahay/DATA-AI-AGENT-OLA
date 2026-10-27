from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "DATA_AGENT_PRESENTATION_BRIEF.docx"
INK = RGBColor(35, 48, 55)
TEAL = RGBColor(0, 112, 116)
CORAL = RGBColor(190, 83, 58)
MUTED = RGBColor(91, 105, 111)


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell(cell, text, *, bold=False, color=INK, size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(1)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None, *, font_size=8.2):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        set_cell(table.rows[0].cells[index], header, bold=True, color=RGBColor(255, 255, 255), size=8.5)
        shade_cell(table.rows[0].cells[index], "006F73")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell(cells[index], value, size=font_size)
            if row_index % 2 == 1:
                shade_cell(cells[index], "F0F5F4")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_callout(document, label, text, fill="E7F1EF", label_color=TEAL):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}  ")
    label_run.bold = True
    label_run.font.color.rgb = label_color
    paragraph.add_run(text)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    for style_name, size, color in (
        ("Title", 30, INK),
        ("Heading 1", 19, INK),
        ("Heading 2", 13, TEAL),
        ("Heading 3", 10.5, INK),
    ):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("DATA AGENT  /  PRESENTATION BRIEF")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"Repository-grounded briefing  |  {date(2026, 10, 28).strftime('%d %B %Y')}")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED


def build_document():
    document = Document()
    configure(document)

    eyebrow = document.add_paragraph()
    eyebrow.paragraph_format.space_before = Pt(36)
    eyebrow.paragraph_format.space_after = Pt(8)
    run = eyebrow.add_run("PRODUCT AND ORGANIZATIONAL USE  /  01")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = TEAL

    title = document.add_paragraph(style="Title")
    title.paragraph_format.space_after = Pt(7)
    title.add_run("DATA AGENT")
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("What it does today, where it can help, and how to grow it responsibly")
    run.font.size = Pt(15)
    run.font.color.rgb = MUTED

    add_callout(
        document,
        "Bottom line",
        "This repository is a promising LangGraph prototype: it routes questions to SQL or ETL workflows over a rideshare-style sample dataset. The useful presentation story is the opportunity to make governed data work more accessible, while being candid that the current execution and access controls need hardening before organizational deployment.",
    )

    add_heading(document, "Executive Summary")
    document.add_paragraph(
        "The Data Agent combines a language-model router with two specialist paths. SQL Analyst translates a business question into PostgreSQL, consults database schema and sample rows, asks a second model to judge whether the query appears safe, and summarizes the returned rows. ETL Analyst can request API extraction or a Pandas transformation through tools. The repository also contains a CSV-to-PostgreSQL loader and five rideshare-style tables: users, vehicles, rides, payments, and ratings."
    )
    document.add_paragraph(
        "At organizational scale, the strongest near-term value is faster self-service exploration and repeatable, governed preparation of data for analysis. Those are potential outcomes, not measured results in this repository. The code shows the workflow shape, but it does not yet demonstrate enterprise permissions, reliable production ETL, audit controls, monitoring, or a tested application entry point."
    )

    add_heading(document, "1. What the Repository Actually Contains")
    add_table(
        document,
        ["Component", "Current behavior", "Evidence / status"],
        [
            ("Data Agent router", "Classifies the latest message as SQL or ETL and invokes one specialist graph.", "Implemented in agents/data_agent.py; no third route or robust fallback."),
            ("SQL Analyst", "Curates the question, fetches public-schema metadata and sample rows, generates SQL, uses an LLM Yes/No judge, executes, then summarizes.", "Implemented in agents/sql_analyst.py and utils/database.py. One checked-in SQL response artifact is an example, not a reliability test."),
            ("ETL Analyst", "Uses model-selected tools to extract from an HTTP endpoint or preview a local file and request Pandas transformation code.", "Implemented in agents/etl_analyst.py and utils/etl_tools.py. Arbitrary generated Python is executed with exec."),
            ("Sample data and loader", "Defines and loads users, vehicles, rides, payments, and ratings tables from CSV files.", "feed_db.py truncates the target tables before loading; running it is destructive."),
            ("Application readiness", "Graph definitions and demonstrations exist, but main.py is empty.", "No production service/API, authentication layer, scheduled orchestration, or automated test suite was found in the repository inventory."),
        ],
        [1.2, 3.65, 3.1],
        font_size=7.8,
    )
    add_callout(
        document,
        "Evidence discipline",
        "A checked-in response file demonstrates one SQL example at some point. The current extract/transform output directories are empty, so this review does not treat the ETL examples as verified successful runs. The repository does not establish production readiness or quantify business impact.",
        "FFF2DE",
        CORAL,
    )

    add_heading(document, "2. How SQL Analyst Helps With Database Access")
    document.add_paragraph(
        "With authorized access to a database, SQL Analyst can act as a natural-language interface for approved analytical questions. A business user might ask, “Which vehicle makes have the most active vehicles?” or “How did completed rides and average fare change month over month?” The intended value is to reduce the time spent translating every question into a first-draft query, not to replace data owners or approve business decisions."
    )
    add_heading(document, "Current query path", 2)
    for item in (
        "The router sends a question to the SQL graph.",
        "The graph rewrites the question and asks DatabaseUtil for table/column metadata plus sample rows from the public schema.",
        "A model generates PostgreSQL; another model returns a structured Yes/No safety judgment.",
        "A Yes result currently leads to direct execution; returned rows are passed to a model for a plain-language response.",
    ):
        add_number(document, item)
    add_heading(document, "Organizational value if governed", 2)
    add_bullet(document, "Self-service answers for recurring operational questions, with analysts spending less time on routine query drafting and more time on interpretation and validation.")
    add_bullet(document, "A common conversational front door to approved metrics and datasets across operations, finance, customer experience, and fleet teams.")
    add_bullet(document, "Faster exploration of joins already represented by the schema, such as rides with payments, riders/drivers, vehicle assignments, and ratings.")
    add_bullet(document, "A path to reusable query templates, reviewable SQL, and cited result context, so common questions become more consistent over time.")
    add_callout(
        document,
        "Database access is not a safety control",
        "The current LLM judge is not a security boundary: utils/database.py executes the query text directly and commits afterward, without deterministic single-SELECT enforcement, a read-only role, or database-enforced row limits. Schema inspection also selects sample rows, which can expose fields such as email and phone to the model. Start with a dedicated least-privilege read-only account, approved schemas/views, minimized metadata, query parsing, timeouts, row caps, and audit logs.",
        "FCEAE4",
        CORAL,
    )

    add_heading(document, "3. How ETL Analyst Helps at Organizational Scale")
    document.add_paragraph(
        "ETL means extract, transform, and load: bringing data from a source, shaping it into an agreed structure, and storing or publishing the result. Across an organization, a governed ETL assistant could help teams onboard partner/API feeds, standardize recurring files, apply approved cleanup rules, and prepare trusted datasets for reporting. It can reduce repetitive handoffs only when transformations are reviewed, repeatable, observable, and owned by the relevant data team."
    )
    add_table(
        document,
        ["Team / workflow", "Possible organizational use", "Guardrail / outcome"],
        [
            ("Operations", "Normalize daily partner or service exports before operational reporting.", "Schema contract, freshness checks, reject/quarantine bad batches."),
            ("Finance", "Prepare payment and settlement files for reconciliation.", "Totals and control-count checks; preserve source and reconciliation trail."),
            ("Data engineering", "Scaffold ingestion for approved APIs and recurring file feeds.", "Versioned pipelines, credentials in a secret manager, retries and lineage."),
            ("Analytics", "Create consistent analysis-ready tables from approved sources.", "Named business definitions, owner review, and documented transformations."),
            ("Customer experience", "Combine ride outcomes and ratings for service-quality reporting.", "Minimize personal data and apply retention and access policies."),
        ],
        [1.35, 3.25, 3.35],
        font_size=7.8,
    )
    etl_note = document.add_paragraph(
        "In this repository, extraction expects a JSON response with a results collection and can write CSV, JSON Lines, or Parquet. Transformation currently previews up to three rows, asks a model for Pandas code, and executes that code. The utility does not itself persist a transformed output; the generated code is expected to do so. The current API request has no timeout, and the output message can sound successful even if code execution returns an error string. These are prototype limitations to fix before scheduled or business-critical use."
    )
    etl_note.paragraph_format.keep_together = True

    add_heading(document, "4. Pros, Challenges, and How to Address Them")
    add_table(
        document,
        ["Area", "Potential benefit / current challenge", "Practical mitigation"],
        [
            ("Accessibility", "Plain-language questions can lower the barrier to first-pass data exploration.", "Show generated SQL and source/metric context; let users inspect, revise, and provide feedback."),
            ("Speed", "Automating repetitive query and file-preparation steps may shorten turnaround; no timing benchmark exists yet.", "Measure baseline versus pilot task time and correctness; keep human review for high-impact outputs."),
            ("SQL security", "A model judge is probabilistic and generated SQL currently reaches direct execution.", "Use read-only database credentials, SQL parsing/allowlisting, approved views, transaction read-only mode, statement timeout, row caps, and query audit."),
            ("Sensitive data", "Schema sampling currently fetches actual rows, potentially including personal fields.", "Give the model metadata only where possible; mask or omit sensitive columns and enforce row/column entitlements in the database."),
            ("ETL execution", "LLM-generated Python runs through unrestricted exec in the application process.", "Replace code generation with allowlisted transformation operations or a declarative spec; isolate workers with no secrets, restricted network, and CPU/memory/time limits."),
            ("Reliability", "API response shape, file formats, generated code, and tool errors are not validated end to end.", "Define input/output schemas, timeouts, retries, idempotency, quarantine, explicit error states, and unit/integration tests."),
            ("Governance", "The repository has no identity, approval, lineage, or audit workflow.", "Integrate SSO/RBAC, dataset ownership, approvals for publishing, immutable run logs, retention rules, and data classification."),
            ("Operations", "No service entry point, metrics, alerting, or retry orchestration is present.", "Deploy behind an authenticated API/job orchestrator; add tracing, run IDs, cost/latency dashboards, alerts, and rollback."),
            ("Dependencies and secrets", "Directly imported runtime packages are missing from requirements.txt; one database credential is hard-coded in source.", "Declare direct dependencies; remove and rotate the exposed credential, then use a managed secret store and secret scanning."),
            ("Trust in answers", "A concise generated summary can conceal ambiguity or a faulty query/result.", "Return SQL, filters, time window, row count, freshness, and source references; validate key calculations against test cases."),
        ],
        [1.15, 3.15, 3.7],
        font_size=7.4,
    )

    add_heading(document, "5. Seven Additional Agent Scopes")
    document.add_paragraph(
        "These are expansion options, not capabilities currently implemented. A sensible sequence is to add bounded, reviewable agents that improve trust and observability before adding autonomous decision-making."
    )
    add_table(
        document,
        ["Proposed agent", "What it could do", "Why it matters / key boundary"],
        [
            ("Data Quality Agent", "Check nulls, duplicates, ranges, referential integrity, and freshness; produce a quality report.", "Makes bad inputs visible before reports or downstream loads. Start read-only; alert or quarantine, never silently rewrite source data."),
            ("Data Catalog and Lineage Agent", "Find approved datasets, owners, definitions, and upstream/downstream lineage.", "Reduces discovery friction and helps assess impact. Use catalog metadata and permission-aware search."),
            ("Business Metrics Agent", "Answer questions using a governed metric catalog (for example completed rides or net revenue).", "Reduces metric drift. Resolve terms to certified definitions; show formula, filters, and owner."),
            ("Anomaly Monitoring Agent", "Explain unusual changes in ride volume, payment failures, cancellations, or ratings.", "Helps teams investigate earlier. Treat detections as signals, include comparison windows, and route alerts to a human owner."),
            ("Privacy and Governance Agent", "Classify columns, flag sensitive-data use, and check requests against policy.", "Improves privacy review. It supplements, never replaces, database authorization or formal compliance review."),
            ("Forecasting and Scenario Agent", "Run approved demand, fleet, or payment forecasting workflows and compare scenarios.", "Supports planning. Use validated statistical/ML models, uncertainty ranges, versioned data, and explicit assumptions."),
            ("Report and Narrative Agent", "Generate recurring KPI summaries, charts, and draft business updates from certified outputs.", "Cuts report assembly effort. Require citations to query/run IDs, freshness labels, and human approval before external distribution."),
        ],
        [1.45, 3.15, 3.4],
        font_size=7.5,
    )
    add_callout(
        document,
        "Suggested first additions",
        "Build Data Quality and Business Metrics first. They directly strengthen the trustworthiness of existing SQL/ETL workflows and create reusable contracts for the later catalog, monitoring, and reporting agents.",
    )

    add_heading(document, "6. Presentation Outline")
    add_table(
        document,
        ["Slide", "Headline", "What to show / say"],
        [
            ("1", "Data Agent: a conversational entry point to data work", "Set the vision; state clearly that the current system is a prototype."),
            ("2", "The workflow has two specialist paths", "Show router -> SQL Analyst or ETL Analyst -> response; explain current scope."),
            ("3", "SQL Analyst can shorten the path from question to query", "Walk through an approved business question, generated SQL, result, and summary; label the example as a prototype demonstration."),
            ("4", "ETL Analyst can reduce repetitive preparation work", "Show an example source-to-clean-output workflow and organizational team use cases; distinguish current limitations from target behavior."),
            ("5", "The rideshare dataset makes the prototype concrete", "Show the five tables and relationships: users, vehicles, rides, payments, ratings."),
            ("6", "Value should be measured, not assumed", "Propose pilot measures: time-to-answer, correctness, analyst effort, data quality, adoption, and cost per task."),
            ("7", "Trust and access controls come before broad rollout", "Present read-only SQL, least privilege, privacy, deterministic validation, isolated ETL, and human review."),
            ("8", "Seven extensions turn a demo into a data platform assistant", "Prioritize Data Quality and Business Metrics; show the other proposed scopes as a phased roadmap."),
            ("9", "Pilot, learn, then expand", "Choose one team and non-sensitive data; define owners, success thresholds, review cadence, and rollback."),
            ("10", "Decision requested", "Ask for a bounded pilot, a data owner, a security reviewer, and time to implement the minimum controls."),
        ],
        [0.55, 2.45, 5.0],
        font_size=7.6,
    )

    add_heading(document, "7. Recommended Phased Roadmap")
    add_table(
        document,
        ["Phase", "Focus", "Exit criteria"],
        [
            ("0. Secure the prototype", "Remove/rotate embedded credentials; make imports side-effect free; eliminate unrestricted exec; make database access read-only and deterministic.", "Security review passes; no unsafe execution path; no destructive database action on import or ordinary run."),
            ("1. Controlled pilot", "Choose one team, one approved dataset, a small set of common questions and transformations; add tests, logs, owner review, and feedback.", "Correctness and time-to-answer meet agreed thresholds; errors and access are auditable; users can identify answer sources."),
            ("2. Establish shared contracts", "Add certified metrics, data quality checks, input/output schemas, and repeatable job execution.", "Definitions and transformation versions have owners; failed runs are visible and recoverable."),
            ("3. Expand by measured demand", "Add catalog/lineage, anomaly monitoring, privacy checks, forecasting, or reporting in response to pilot evidence.", "Each added agent has a named owner, bounded tools, evaluation set, permission model, and operational SLO."),
        ],
        [1.15, 4.1, 2.75],
        font_size=7.6,
    )

    add_heading(document, "Pilot Measures")
    add_bullet(document, "Answer quality: percentage of sampled questions whose SQL, result, and interpretation pass data-owner review.")
    add_bullet(document, "Time-to-answer: median end-to-end time compared with the current workflow for the same tasks.")
    add_bullet(document, "ETL reliability: successful validated runs, schema drift incidents, retries, and recoverable failures.")
    add_bullet(document, "Safety: denied unauthorized requests, policy violations, sensitive-data exposure incidents, and audit completeness.")
    add_bullet(document, "Adoption and cost: repeat usage by intended users, user corrections, latency, and model/infrastructure cost per completed task.")

    add_heading(document, "8. Speaker Notes and Closing Message")
    add_callout(
        document,
        "Closing message",
        "The opportunity is not simply to let a model query a database. It is to make approved data easier to discover, prepare, and explain while preserving ownership, access control, and evidence. This codebase is a useful prototype for that direction; a carefully measured, least-privilege pilot is the credible next step.",
    )
    document.add_paragraph(
        "Suggested opening: “Our teams spend time finding data, translating questions into queries, and repeating file preparation. This prototype explores one conversational entry point for those tasks. Today I’ll show what it does, where it could help, and the controls we need before using it with organizational data.”"
    )

    add_heading(document, "Repository References")
    add_table(
        document,
        ["File", "Relevance"],
        [
            ("agents/data_agent.py", "SQL/ETL classification and graph routing."),
            ("agents/sql_analyst.py", "Question curation, schema prompt, SQL generation, LLM judge, execution path, response generation."),
            ("agents/etl_analyst.py", "Extraction/transformation tools and tool-calling graph."),
            ("utils/database.py", "Schema/sample-row retrieval and SQL execution; import-time database access and embedded credential."),
            ("utils/etl_tools.py", "HTTP extraction, file preview, path resolution, unrestricted Python execution."),
            ("Models/schema.py", "Pydantic state and router contracts."),
            ("feed_db.py", "Five-table PostgreSQL schema, indexes, CSV loader, destructive truncate."),
            ("requirements.txt", "Declared runtime dependencies; compare with direct imports before deployment."),
        ],
        [2.4, 5.6],
        font_size=8,
    )
    document.add_paragraph(
        "Prepared from the checked-in source and artifacts in this workspace on 28 October 2026. This document does not connect to a live database, inspect environment secrets, or claim that proposed capabilities are implemented."
    )

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()